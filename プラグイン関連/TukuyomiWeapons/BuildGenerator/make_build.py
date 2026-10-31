#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ビルド申請のExcelから、tukuyomi_build.sk 用の「名前→BuildID」と「BuildIDの中身」の行を作る。

使い方
  python make_build.py 申請.xlsx                       ... build_registry.txt を作る(貼り付け用)
  python make_build.py 申請.xlsx tukuyomi_build.sk     ... tukuyomi_build.sk の自動生成範囲を直接書き換える(.bakを残す)

必要なもの: Python 3.8以上 と openpyxl (pip install openpyxl)
"""
import sys, os, re, shutil
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("openpyxl が入っていません。次を実行してください:  pip install openpyxl")
    sys.exit(1)

# ============================================================
#  対応表 (新しい武器やアイテムが増えたら、ここに足す)
# ============================================================
# カスタム武器(メイン/サブ): Excelの表示名(コストを除く) → 武器ID
CUSTOM = {
    "スーパーノヴァ": "light_sword",
    "鋸ナタ": "saw_cleaver",
    "拒絶": "rejection",
    "獣肉断ち": "beast_cleaver",
    "パイルハンマー": "pile_hammer",
    "仕込み杖": "cane_sheathed",
    "FUSHI爆弾": "fushi_bomb",
    "もと光るダムダム弾": "timed_bomb",
    "自律稼働無人爆弾「月人」": "tsukibito_bomb",
    "彩葉特製、毒素マシマシパンケーキ": "pancake",
    "ギター「共振波動」": "guitar_resonance",
    "ガンランス「MIKA-D0」": "gunlance_mika_d0",
    "狩猟笛": "hunting_horn",
}
# ウルト: Excelの表示名 → 武器ID
ULT = {
    "光の剣": "light_sword_ult",
    "古狩人の遺骨": "old_hunter_bone",
    "月光の聖剣": "moonlight_holy_sword",
    "彼方への呼びかけ": "call_to_beyond",
    "パイルバンカー": "pile_bunker",
    "神の怒り": "wrath_of_god",
    "竜撃砲": "dragon_cannon",
    "共鳴震": "resonance_quake",
    "音撃震": "sonic_quake",
    "8000年の追憶「弱」": "reminiscence_8000_weak",
    "唐傘弾": "karakasa_shell",
    "星界からの使者": "envoy_from_stars",
    "大失態": "big_blunder",
}
# 防具: 「○○一式」の○○ → 素材
ARMOR = {
    "革": "leather", "金": "golden", "チェーン": "chainmail",
    "鉄": "iron", "ダイヤ": "diamond", "ネザライト": "netherite",
}
# バニラ武器: Excelの表示名 → Skriptのアイテム名
VANILLA = {
    "木の剣": "wooden sword", "石の剣": "stone sword", "鉄の剣": "iron sword",
    "ダイヤの剣": "diamond sword", "ネザライトの剣": "netherite sword",
    "木の斧": "wooden axe", "石の斧": "stone axe", "鉄の斧": "iron axe",
    "ダイヤの斧": "diamond axe", "ネザライトの斧": "netherite axe",
    "弓": "bow", "クロスボウ": "crossbow", "トライデント": "trident", "メイス": "mace",
}
ARROW_ITEM = "vanilla:arrow:64"   # 矢「使用する」のときに渡す内容

SHEET = "フォームの回答 1"
# 列の位置(0始まり)。Excelの並びを変えたらここを直す
C_TIME, C_NAME, C_ARMOR, C_SHIELD, C_ELYTRA, C_VANILLA, C_ARROW, C_MAIN, C_SUB = 0, 1, 2, 3, 4, 5, 6, 7, 8
C_COST, C_LIMITOK, C_ULT, C_ID = 17, 19, 20, 21     # 合計コスト / 判定 / ウルト / BuildID (数式の結果)

START = "# >>> GENERATED-BUILDS START"
END = "# <<< GENERATED-BUILDS END"


def s(v):
    return "" if v is None else str(v).replace("\u200b", "").strip()


def strip_cost(text):
    """「鋸ナタ-250」→「鋸ナタ」"""
    return re.sub(r"\s*-\s*\d+(\.\d+)?$", "", s(text)).strip()


def split_multi(text):
    """複数選択のセルを分ける(区切りは半角カンマと改行。名前の中の「、」では分けない)"""
    return [p.strip() for p in re.split(r",\s*|\n", s(text)) if p.strip()]


def is_none(text):
    return s(text) in ("", "なし", "使用しない", "無し")


def build_row(row, rownum, errors, ult_table):
    items = []
    notes = {}

    # 防具
    a = strip_cost(row[C_ARMOR])
    if not is_none(a):
        key = a.replace("一式", "")
        if key in ARMOR:
            items.append("armor:" + ARMOR[key])
        else:
            errors.append(f"{rownum}行目: 未対応の防具「{a}」")
    notes["防具"] = a or "なし"

    # 盾・エリトラ
    if s(row[C_SHIELD]).startswith("有り"):
        items.append("vanilla:shield")
    if s(row[C_ELYTRA]).startswith("有り"):
        items.append("vanilla:elytra")

    # バニラ武器
    van = []
    for part in split_multi(row[C_VANILLA]):
        name = strip_cost(part)
        if name in VANILLA:
            items.append("vanilla:" + VANILLA[name])
            van.append(name)
        else:
            errors.append(f"{rownum}行目: 未対応のバニラ武器「{name}」")
    notes["バニラ"] = "、".join(van) or "なし"

    # 矢
    if s(row[C_ARROW]).startswith("使用する"):
        items.append(ARROW_ITEM)
    notes["矢"] = "あり" if s(row[C_ARROW]).startswith("使用する") else "なし"

    # メイン + ウルト
    main = strip_cost(row[C_MAIN])
    ult = s(row[C_ULT])
    if not ult and main in ult_table:
        ult = ult_table[main]        # 数式の結果が保存されていないときは、「ウルト対応表」から引く
    if not is_none(main):
        if main in CUSTOM:
            items.append("custom:" + CUSTOM[main])
        else:
            errors.append(f"{rownum}行目: 未対応のメイン武器「{main}」")
        if ult:
            if ult in ULT:
                items.append("custom:" + ULT[ult])
            else:
                errors.append(f"{rownum}行目: 未対応のウルト「{ult}」")
        else:
            errors.append(f"{rownum}行目: メイン「{main}」のウルトが空です(Excelの数式の結果が入っていません)")
    notes["メイン"] = main or "なし"
    notes["ウルト"] = ult or "なし"

    # サブ
    subs = []
    for part in split_multi(row[C_SUB]):
        name = strip_cost(part)
        if is_none(name):
            continue
        if name in CUSTOM:
            items.append("custom:" + CUSTOM[name])
            subs.append(name)
        else:
            errors.append(f"{rownum}行目: 未対応のサブ武器「{name}」")
    notes["サブ"] = "、".join(subs) or "なし"

    # 重複を除く(順番は保つ)
    seen, uniq = set(), []
    for it in items:
        if it not in seen:
            seen.add(it)
            uniq.append(it)
    return uniq, notes


def skript_list(items):
    q = [f'"{i}"' for i in items]
    if len(q) <= 1:
        return ", ".join(q)
    return ", ".join(q[:-1]) + " and " + q[-1]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    xlsx = next((a for a in args if a.lower().endswith((".xlsx", ".xlsm"))), None)
    sk = next((a for a in args if a.lower().endswith(".sk")), None)
    if not xlsx:
        print(__doc__)
        return 1

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]

    # 「ウルト対応表」シート: メイン武器 → ウルト
    ult_table = {}
    if "ウルト対応表" in wb.sheetnames:
        for r in wb["ウルト対応表"].iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[1]:
                ult_table[strip_cost(r[0])] = s(r[1])

    errors, warnings = [], []
    builds = []   # (rownum, name, key, id, items, notes, cost, ok, ts)
    for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = list(row) + [None] * (24 - len(row))
        name = s(row[C_NAME])
        if not name:
            continue
        key = name.lower().replace(" ", "_")
        if " " in name:
            warnings.append(f"{rownum}行目: 名前「{name}」に空白があります。統合版は空白が _ に変わる設定のため _ にしました(実際の名前と合わせて確認してください)")
        if re.search(r"[%{}:]", key):
            errors.append(f"{rownum}行目: 名前「{name}」にSkriptで使えない文字(% {{ }} :)があります")
        bid = s(row[C_ID])
        if not bid.startswith("BUILD-"):
            bid = f"BUILD-{rownum - 1:03d}"      # Excelの数式と同じ規則
        items, notes = build_row(row, rownum, errors, ult_table)
        cost = s(row[C_COST])
        ok = s(row[C_LIMITOK])
        ts = row[C_TIME] if isinstance(row[C_TIME], datetime) else None
        builds.append((rownum, name, key, bid, items, notes, cost, ok, ts))

    if not builds:
        print("ビルドの申請が見つかりませんでした。")
        return 1

    # 同じ名前が複数ある場合は、新しい申請を採用
    owner = {}
    for b in builds:
        k = b[2]
        if k in owner:
            prev = owner[k]
            newer = b if (b[8] and prev[8] and b[8] >= prev[8]) or not (b[8] and prev[8]) else prev
            warnings.append(f"名前「{b[1]}」の申請が複数あります({prev[3]} と {b[3]})。新しいほうの {newer[3]} を採用しました")
            owner[k] = newer
        else:
            owner[k] = b

    # 出力
    lines = []
    lines.append("# --- プレイヤー名(小文字) → BuildID ---")
    lines.append("# 統合版は先頭に「.」を付けた名前で登録")
    for k, b in owner.items():
        lines.append(f'set {{tkb::owner::{k}}} to "{b[3]}"')
    lines.append("")
    for b in builds:
        rownum, name, key, bid, items, notes, cost, ok, ts = b
        head = f"# --- {bid} / {name} ---"
        if cost:
            head += f"  (コスト合計 {cost} / 判定 {ok})"
        lines.append(head)
        lines.append(f"# 防具: {notes['防具']} / メイン: {notes['メイン']} / ウルト: {notes['ウルト']}")
        lines.append(f"# バニラ: {notes['バニラ']} / 矢: {notes['矢']}")
        lines.append(f"# サブ: {notes['サブ']}")
        lines.append(f"set {{tkb::items::{bid}::*}} to {skript_list(items)}")
        lines.append("")
        if ok and ok != "OK":
            warnings.append(f"{bid} / {name}: Excelの判定が「{ok}」です(コスト合計 {cost})")
    while lines and lines[-1] == "":
        lines.pop()

    out_dir = os.path.dirname(os.path.abspath(xlsx))
    out_path = os.path.join(out_dir, "build_registry.txt")
    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")
    print(f"{len(builds)} 件のビルドを読み込みました → {out_path}")

    for w in warnings:
        print("[注意]", w)
    for e in errors:
        print("[エラー]", e)

    if sk:
        if errors:
            print("エラーがあるため、tukuyomi_build.sk は書き換えませんでした。対応表を直してから、もう一度実行してください。")
            return 2
        with open(sk, "r", encoding="utf-8", newline="") as f:
            text = f.read()
        nl = "\r\n" if "\r\n" in text else "\n"
        m = re.search(r"^([ \t]*)" + re.escape(START) + r".*?" + nl + r"(.*?)^[ \t]*" + re.escape(END), text, re.S | re.M)
        if not m:
            print(f"{sk} に目印の行({START} / {END})が見つかりません。")
            return 2
        indent = m.group(1)
        body = nl.join((indent + l) if l else "" for l in lines) + nl
        start_line_end = text.index(nl, m.start()) + len(nl)
        new_text = text[:start_line_end] + body + indent + END + text[m.end():]
        shutil.copyfile(sk, sk + ".bak")
        with open(sk, "w", encoding="utf-8", newline="") as f:
            f.write(new_text)
        print(f"{sk} を書き換えました(元のファイルは {os.path.basename(sk)}.bak)。サーバーで /sk reload tukuyomi_build を実行してください。")
    else:
        print("build_registry.txt の中身を tukuyomi_build.sk の on load: の中に貼り付けてください。")
        print("(.sk も一緒に指定すると、自動で書き換えます)")
    return 2 if errors else 0


if __name__ == "__main__":
    code = main()
    sys.exit(code)
